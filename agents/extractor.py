"""Document discovery, format routing, parsing orchestration, and evidence-grounded structured field extraction."""

from __future__ import annotations

import csv
from datetime import date
import io
import json
import logging
from pathlib import Path
import re
from typing import Any, Callable, Dict, List, Optional, Protocol, Set, Tuple
import xml.etree.ElementTree as ET
import zipfile

import openpyxl

from agents.models import (
    DocumentContent,
    DocumentPage,
    ExtractedField,
    ExtractionResult,
    FailureCategory,
    FailureRecord,
    SourceLocation,
    VerificationStatus,
)
from agents.reliability import (
    OCRQualityError,
    ReliableOCRRunner,
    SimpleGrayscaleImage,
    classify_exception,
)

try:
    from PIL import Image
    HAS_PIL = True
except ImportError:
    Image = None  # type: ignore
    HAS_PIL = False

try:
    import pymupdf
    HAS_PYMUPDF = True
except ImportError:
    pymupdf = None  # type: ignore
    HAS_PYMUPDF = False

try:
    import pypdf
    HAS_PYPDF = True
except ImportError:
    pypdf = None  # type: ignore
    HAS_PYPDF = False

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS: Set[str] = {
    ".pdf",
    ".txt",
    ".json",
    ".csv",
    ".xlsx",
    ".docx",
    ".png",
    ".jpg",
    ".jpeg",
    ".tiff",
    ".bmp",
}


class OCREngine(Protocol):
    """Protocol for OCR engines."""
    def extract_text(self, image: Any) -> str:
        ...


class DefaultOCREngine:
    """Default OCR engine that attempts pytesseract if available, or delegates to a fallback handler."""
    def __init__(self, fallback_handler: Optional[Callable[[Any], str]] = None):
        self.fallback_handler = fallback_handler

    def extract_text(self, image: Any) -> str:
        if self.fallback_handler:
            return self.fallback_handler(image)

        try:
            import pytesseract
            return pytesseract.image_to_string(image)
        except (ImportError, Exception) as exc:
            raise OCRQualityError(
                f"OCR failed: No OCR engine or pytesseract binary available ({exc})"
            ) from exc


# ---------------------------------------------------------------------------
# File Discovery & Candidate Identification
# ---------------------------------------------------------------------------

def discover_files(directory: str | Path, supported_exts: Optional[Set[str]] = None) -> List[Path]:
    """Discover supported document files inside a directory (including subdirectories)."""
    dir_path = Path(directory)
    if not dir_path.is_dir():
        return []
    valid_exts = supported_exts or SUPPORTED_EXTENSIONS
    found: List[Path] = []
    for p in sorted(dir_path.rglob("*")):
        if p.is_file() and p.suffix.lower() in valid_exts:
            found.append(p)
    return found


def normalize_query_terms(query: str) -> List[str]:
    """Extract significant vendor/entity search keywords from an instruction query."""
    stop_words = {
        "find", "and", "read", "the", "invoice", "invoices", "from", "for", "a", "an", "of",
        "to", "in", "document", "file", "please", "get", "show", "extract", "latest", "recent",
        "record", "it", "finance", "system", "tell", "me", "what", "is", "whats", "due",
        "when", "does", "our", "contract", "with", "expire", "expires", "pay", "supplies",
        "enterprises", "corporation", "corp", "inc", "llc", "ltd", "services", "chemicals",
        "logistics", "components", "industrial",
    }
    cleaned = re.sub(r"[^\w\s]", " ", query.lower())
    words = [w.strip() for w in cleaned.split() if w.strip()]
    meaningful = [w for w in words if w not in stop_words]
    if not meaningful:
        # Fallback if only company suffix words remained
        fallback_stop = {"find", "and", "read", "the", "invoice", "from", "for", "in", "record", "latest"}
        meaningful = [w for w in words if w not in fallback_stop]
    return meaningful if meaningful else words


def _resolve_latest_among_candidates(
    candidates: List[Path],
    ocr_engine: Optional[OCREngine] = None,
) -> Tuple[List[Path], bool, Optional[str]]:
    """Resolve the authoritative latest invoice among multiple candidates from the same vendor.

    1. Filters out invoices explicitly superseded by another candidate's 'supersedes' clause.
    2. Checks extracted invoice_date normalization: if any candidate has an ambiguous date,
       refuses to guess and reports ambiguity.
    3. Selects the candidate with the latest unambiguous invoice_date.
    """
    from agents.verification import normalize_date

    parsed_candidates: List[Dict[str, Any]] = []
    superseded_numbers: Set[str] = set()

    for path in candidates:
        try:
            doc = read_document(path, ocr_engine=ocr_engine)
            res = extract_fields(
                doc,
                requested_fields=["invoice_number", "invoice_date", "supersedes"],
            )
            inv_num = res.fields["invoice_number"].extracted_value
            inv_date_raw = res.fields["invoice_date"].extracted_value
            supersedes_val = res.fields["supersedes"].extracted_value

            if supersedes_val:
                superseded_numbers.add(str(supersedes_val).strip().upper())

            norm_date, is_ambiguous, err = normalize_date(inv_date_raw)
            parsed_candidates.append({
                "path": path,
                "invoice_number": str(inv_num).strip().upper() if inv_num else None,
                "inv_date_raw": inv_date_raw,
                "norm_date": norm_date,
                "is_ambiguous": is_ambiguous,
                "error": err,
            })
        except Exception as exc:
            logger.warning("Failed inspecting candidate %s for date resolution: %s", path.name, exc)
            parsed_candidates.append({
                "path": path,
                "invoice_number": None,
                "inv_date_raw": None,
                "norm_date": None,
                "is_ambiguous": True,
                "error": str(exc),
            })

    # Filter out superseded invoices
    active = [
        c for c in parsed_candidates
        if not (c["invoice_number"] and c["invoice_number"] in superseded_numbers)
    ]
    if not active:
        active = parsed_candidates

    if len(active) == 1 and not active[0]["is_ambiguous"]:
        return [active[0]["path"]], False, None

    # If any active candidate has an ambiguous or unparseable invoice date, cannot reliably break tie
    if any(c["is_ambiguous"] or c["norm_date"] is None for c in active):
        return (
            [c["path"] for c in active],
            True,
            f"Ambiguous candidate invoices with unresolved or ambiguous dates: {[c['path'].name for c in active]}",
        )

    # Sort by unambiguous normalized invoice_date descending
    active.sort(key=lambda c: c["norm_date"], reverse=True)
    top_date = active[0]["norm_date"]
    top_tied = [c for c in active if c["norm_date"] == top_date]

    if len(top_tied) > 1:
        return (
            [c["path"] for c in top_tied],
            True,
            f"Ambiguous: multiple invoices share the latest date {top_date}: {[c['path'].name for c in top_tied]}",
        )

    return [top_tied[0]["path"]], False, None


def find_candidate_documents(
    directory: str | Path,
    query: str,
    ocr_engine: Optional[OCREngine] = None,
) -> Tuple[List[Path], bool, Optional[str]]:
    """Identify candidate documents relevant to the query.

    Returns:
        (candidates, is_ambiguous, reason_if_any)
    """
    dir_path = Path(directory)
    if not dir_path.is_dir():
        return [], False, f"Directory not found: {directory}"

    all_files = discover_files(dir_path)
    if not all_files:
        return [], False, f"No supported documents found in directory: {directory}"

    q_lower = query.lower()
    wants_contract = "contract" in q_lower or "msa" in q_lower or "nda" in q_lower or "expire" in q_lower
    wants_policy = "policy" in q_lower
    wants_latest = "latest" in q_lower or "recent" in q_lower or "current" in q_lower

    terms = normalize_query_terms(query)
    if not terms:
        return all_files, len(all_files) > 1, "Ambiguous: no specific company or document search term provided"

    matches: List[Tuple[Path, int]] = []
    for file_path in all_files:
        # Folder filtering if subfolders exist
        parent_name = file_path.parent.name.lower()
        if wants_contract and parent_name == "invoices":
            continue
        if not wants_contract and not wants_policy and parent_name in ("contracts", "policies"):
            continue

        score = 0
        name_lower = file_path.name.lower()
        for term in terms:
            if term in name_lower:
                score += 3

        # Peek content if filename didn't match
        if score == 0:
            try:
                if file_path.suffix.lower() in {".txt", ".json", ".csv"}:
                    snippet = file_path.read_text(encoding="utf-8", errors="ignore")[:2000].lower()
                elif file_path.suffix.lower() == ".pdf" and HAS_PYMUPDF:
                    with pymupdf.open(file_path) as doc_peek:
                        snippet = (doc_peek[0].get_text() if len(doc_peek) > 0 else "").lower()
                else:
                    snippet = ""
                for term in terms:
                    if term in snippet:
                        score += 1
            except Exception as exc:
                logger.debug("Content peek skipped for %s: %s", file_path.name, exc)

        if score > 0:
            matches.append((file_path, score))

    if not matches:
        return [], False, f"No documents matched / no invoice found matching query terms: {terms}"

    matches.sort(key=lambda x: x[1], reverse=True)
    best_score = matches[0][1]
    best_candidates = [f for f, s in matches if s == best_score]

    if len(best_candidates) > 1:
        if wants_latest:
            return _resolve_latest_among_candidates(best_candidates, ocr_engine=ocr_engine)
        return best_candidates, True, (
            f"Ambiguous: multiple matching documents found for query '{query}': "
            f"{[f.name for f in best_candidates]}"
        )

    return best_candidates, False, None


# ---------------------------------------------------------------------------
# Format-Specific Readers
# ---------------------------------------------------------------------------

def read_txt(file_path: Path) -> DocumentContent:
    """Read plain text document."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            text = file_path.read_text(encoding=encoding)
            return DocumentContent(
                file_path=file_path,
                file_format="txt",
                raw_text=text,
                pages=[DocumentPage(page_num=1, text=text, is_ocr=False)],
                metadata={"encoding": encoding},
            )
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Unable to decode text file with supported encodings: {file_path}")


def read_json(file_path: Path) -> DocumentContent:
    """Read JSON document preserving structure and formatted text."""
    try:
        content = file_path.read_text(encoding="utf-8")
        parsed = json.loads(content)
    except Exception as exc:
        raise ValueError(f"Corrupt or invalid JSON file: {exc}") from exc
    formatted_text = json.dumps(parsed, indent=2)
    return DocumentContent(
        file_path=file_path,
        file_format="json",
        raw_text=formatted_text,
        pages=[DocumentPage(page_num=1, text=formatted_text, is_ocr=False)],
        structured_data=parsed,
    )


def read_csv(file_path: Path) -> DocumentContent:
    """Read CSV tabular data preserving row/column references."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            with open(file_path, "r", encoding=encoding, newline="") as f:
                reader = csv.reader(f)
                rows = list(reader)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ValueError(f"Unable to decode CSV file: {file_path}")

    tables: List[Dict[str, Any]] = []
    lines: List[str] = []
    if rows:
        headers = rows[0]
        tables.append({"headers": headers, "rows": rows[1:]})
        for idx, row in enumerate(rows, start=1):
            lines.append(f"Row {idx}: " + ", ".join(f"{h}={val}" for h, val in zip(headers, row) if val))

    raw_text = "\n".join(lines)
    return DocumentContent(
        file_path=file_path,
        file_format="csv",
        raw_text=raw_text,
        pages=[DocumentPage(page_num=1, text=raw_text, is_ocr=False)],
        tables=tables,
    )


def read_xlsx(file_path: Path) -> DocumentContent:
    """Read Excel spreadsheet preserving sheet names, headers, and exact cell coordinates."""
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
    except Exception as exc:
        raise ValueError(f"Corrupt or invalid XLSX file: {exc}") from exc

    tables: List[Dict[str, Any]] = []
    sheet_lines: List[str] = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        sheet_rows: List[List[Any]] = []
        sheet_coords: List[List[str]] = []
        cell_entries: List[str] = []

        for row in sheet.iter_rows(values_only=False):
            row_vals = []
            row_c = []
            for cell in row:
                val = cell.value
                row_vals.append(val)
                row_c.append(cell.coordinate)
                if val is not None and str(val).strip():
                    coord = f"{sheet_name}!{cell.coordinate}"
                    cell_entries.append(f"[{coord}] {val}")
            if any(v is not None for v in row_vals):
                sheet_rows.append(row_vals)
                sheet_coords.append(row_c)

        if sheet_rows:
            headers = [str(v) if v is not None else "" for v in sheet_rows[0]]
            tables.append({
                "sheet": sheet_name,
                "headers": headers,
                "rows": sheet_rows[1:],
                "all_rows": sheet_rows,
                "all_coords": sheet_coords,
            })

        if cell_entries:
            sheet_lines.append(f"Sheet '{sheet_name}':\n" + "\n".join(cell_entries))

    raw_text = "\n\n".join(sheet_lines)
    return DocumentContent(
        file_path=file_path,
        file_format="xlsx",
        raw_text=raw_text,
        pages=[DocumentPage(page_num=1, text=raw_text, is_ocr=False)],
        tables=tables,
    )


def read_docx(file_path: Path) -> DocumentContent:
    """Read Word (.docx) document using zipfile and XML ElementTree."""
    try:
        with zipfile.ZipFile(file_path) as z:
            xml_content = z.read("word/document.xml")
    except Exception as exc:
        raise ValueError(f"Corrupt or invalid DOCX file: {exc}") from exc

    tree = ET.fromstring(xml_content)
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}

    paragraphs: List[str] = []
    for p in tree.iter(f"{{{ns['w']}}}p"):
        texts = [t.text for t in p.iter(f"{{{ns['w']}}}t") if t.text]
        if texts:
            paragraphs.append("".join(texts))

    raw_text = "\n".join(paragraphs)
    return DocumentContent(
        file_path=file_path,
        file_format="docx",
        raw_text=raw_text,
        pages=[DocumentPage(page_num=1, text=raw_text, is_ocr=False)],
    )


def read_image(file_path: Path, ocr_engine: Optional[OCREngine] = None) -> DocumentContent:
    """Read image file using OCR with bounded reliability."""
    engine = ocr_engine or DefaultOCREngine()
    reliable_engine = engine if isinstance(engine, ReliableOCRRunner) else ReliableOCRRunner(engine)

    image_obj: Any = None
    img_meta: Dict[str, Any] = {}

    if HAS_PIL and Image is not None:
        try:
            with Image.open(file_path) as img:
                image_obj = img.copy()
                img_meta = {"image_format": img.format, "size": img.size}
        except Exception as exc:
            raise ValueError(f"Corrupt or invalid image file: {exc}") from exc
    elif HAS_PYMUPDF and pymupdf is not None:
        try:
            pix = pymupdf.Pixmap(str(file_path))
            image_obj = SimpleGrayscaleImage(pix.width, pix.height, pix.samples)
            img_meta = {"image_format": file_path.suffix.lstrip(".").upper(), "size": (pix.width, pix.height)}
        except Exception as exc:
            raise ValueError(f"Corrupt or invalid image file: {exc}") from exc
    else:
        raise ValueError("No image reading library (Pillow or PyMuPDF) available")

    text = reliable_engine.extract_text(image_obj)
    return DocumentContent(
        file_path=file_path,
        file_format="image",
        raw_text=text,
        pages=[DocumentPage(page_num=1, text=text, is_ocr=True, image_count=1)],
        metadata={**img_meta, "recovery_log": getattr(reliable_engine, "last_recovery_log", [])},
    )


def _is_extracted_pdf_text_sufficient(text: str, image_count: int) -> bool:
    """Determine whether direct PDF text is sufficient rather than merely non-empty."""
    cleaned = text.strip()
    if len(cleaned) < 15:
        return False
    alnum = sum(1 for c in cleaned if c.isalnum())
    if alnum < 10:
        return False
    # If page has embedded images and only a tiny header/footer (< 25 chars), treat as scanned
    if image_count > 0 and len(cleaned) < 25:
        return False
    return True


def read_pdf(file_path: Path, ocr_engine: Optional[OCREngine] = None) -> DocumentContent:
    """Read PDF document: attempt direct text extraction first; use OCR on rendered/embedded page images for scanned pages."""
    pages: List[DocumentPage] = []
    full_text_parts: List[str] = []
    recovery_logs: List[str] = []

    if HAS_PYMUPDF and pymupdf is not None:
        try:
            doc = pymupdf.open(file_path)
            total_pages = len(doc)
        except Exception as exc:
            raise ValueError(f"Corrupt or invalid PDF file: {exc}") from exc

        try:
            for idx, page in enumerate(doc, start=1):
                extracted_text = (page.get_text() or "").strip()
                img_list = page.get_images()
                num_images = len(img_list)

                if _is_extracted_pdf_text_sufficient(extracted_text, num_images):
                    pages.append(
                        DocumentPage(
                            page_num=idx,
                            text=extracted_text,
                            is_ocr=False,
                            image_count=num_images,
                        )
                    )
                    full_text_parts.append(f"--- Page {idx} ---\n{extracted_text}")
                else:
                    raw_engine = ocr_engine or DefaultOCREngine()
                    engine = raw_engine if isinstance(raw_engine, ReliableOCRRunner) else ReliableOCRRunner(raw_engine)
                    page_images: List[Any] = []
                    ocr_texts: List[str] = []

                    # Render page via PyMuPDF pixmap so even complex/inline scanned pages can be OCR'd
                    pix = page.get_pixmap()
                    if HAS_PIL and Image is not None:
                        mode = "RGBA" if pix.alpha else "RGB"
                        pil_img = Image.frombytes(mode, (pix.width, pix.height), pix.samples)
                        page_images.append(pil_img)
                        ocr_res = engine.extract_text(pil_img)
                    else:
                        simple_img = SimpleGrayscaleImage(pix.width, pix.height, pix.samples)
                        page_images.append(simple_img)
                        ocr_res = engine.extract_text(simple_img)

                    if ocr_res.strip():
                        ocr_texts.append(ocr_res.strip())
                    if hasattr(engine, "last_recovery_log"):
                        recovery_logs.extend(engine.last_recovery_log)

                    page_ocr_text = "\n".join(ocr_texts) if ocr_texts else extracted_text
                    pages.append(
                        DocumentPage(
                            page_num=idx,
                            text=page_ocr_text,
                            is_ocr=True,
                            image_count=max(num_images, len(page_images)),
                            images=page_images,
                        )
                    )
                    full_text_parts.append(f"--- Page {idx} (OCR) ---\n{page_ocr_text}")
        finally:
            doc.close()

        return DocumentContent(
            file_path=file_path,
            file_format="pdf",
            raw_text="\n\n".join(full_text_parts),
            pages=pages,
            metadata={"total_pages": total_pages, "recovery_log": recovery_logs},
        )

    if HAS_PYPDF and pypdf is not None:
        try:
            reader = pypdf.PdfReader(file_path)
        except Exception as exc:
            raise ValueError(f"Corrupt or invalid PDF file: {exc}") from exc

        for idx, page in enumerate(reader.pages, start=1):
            extracted_text = (page.extract_text() or "").strip()
            num_images = len(page.images) if hasattr(page, "images") else 0

            if _is_extracted_pdf_text_sufficient(extracted_text, num_images):
                pages.append(
                    DocumentPage(
                        page_num=idx,
                        text=extracted_text,
                        is_ocr=False,
                        image_count=num_images,
                    )
                )
                full_text_parts.append(f"--- Page {idx} ---\n{extracted_text}")
            else:
                raw_engine = ocr_engine or DefaultOCREngine()
                engine = raw_engine if isinstance(raw_engine, ReliableOCRRunner) else ReliableOCRRunner(raw_engine)
                ocr_texts = []
                page_images = []
                if hasattr(page, "images") and page.images:
                    for img_file in page.images:
                        if hasattr(img_file, "image") and img_file.image is not None:
                            pil_img = img_file.image
                        elif HAS_PIL and Image is not None:
                            pil_img = Image.open(io.BytesIO(img_file.data))
                        else:
                            pil_img = SimpleGrayscaleImage(100, 100, img_file.data)
                        page_images.append(pil_img)
                        ocr_result = engine.extract_text(pil_img)
                        if ocr_result.strip():
                            ocr_texts.append(ocr_result.strip())
                        if hasattr(engine, "last_recovery_log"):
                            recovery_logs.extend(engine.last_recovery_log)

                page_ocr_text = "\n".join(ocr_texts) if ocr_texts else extracted_text
                pages.append(
                    DocumentPage(
                        page_num=idx,
                        text=page_ocr_text,
                        is_ocr=True,
                        image_count=len(page_images),
                        images=page_images,
                    )
                )
                full_text_parts.append(f"--- Page {idx} (OCR) ---\n{page_ocr_text}")

        return DocumentContent(
            file_path=file_path,
            file_format="pdf",
            raw_text="\n\n".join(full_text_parts),
            pages=pages,
            metadata={"total_pages": len(reader.pages), "recovery_log": recovery_logs},
        )

    raise ValueError("No PDF library (PyMuPDF or pypdf) is installed")


# ---------------------------------------------------------------------------
# Router
# ---------------------------------------------------------------------------

def read_document(file_path: str | Path, ocr_engine: Optional[OCREngine] = None) -> DocumentContent:
    """Route document to format-specific reader based on file extension and validate."""
    path = Path(file_path)
    if not path.is_file():
        raise FileNotFoundError(f"File does not exist: {path}")

    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported format '{ext}' for file {path.name}")

    if ext == ".txt":
        return read_txt(path)
    elif ext == ".json":
        return read_json(path)
    elif ext == ".csv":
        return read_csv(path)
    elif ext == ".xlsx":
        return read_xlsx(path)
    elif ext == ".docx":
        return read_docx(path)
    elif ext in {".png", ".jpg", ".jpeg", ".tiff", ".bmp"}:
        return read_image(path, ocr_engine=ocr_engine)
    elif ext == ".pdf":
        return read_pdf(path, ocr_engine=ocr_engine)
    else:
        raise ValueError(f"Unsupported format '{ext}' for file {path.name}")


# ---------------------------------------------------------------------------
# Structured Field Extraction
# ---------------------------------------------------------------------------

DEFAULT_REQUESTED_FIELDS: List[str] = [
    "vendor",
    "invoice_number",
    "invoice_date",
    "due_date",
    "total_amount",
    "currency",
    "subtotal",
    "tax",
]

FIELD_KEY_SYNONYMS: Dict[str, List[str]] = {
    "vendor": ["vendor", "vendor_name", "company", "company_name", "issuer", "seller", "supplier"],
    "invoice_number": ["invoice_number", "invoice_no", "invoice_id", "inv_num", "inv_no", "number", "id"],
    "invoice_date": ["invoice_date", "date", "bill_date", "issue_date", "dated"],
    "due_date": ["due_date", "payment_due", "pay_by", "due"],
    "total_amount": ["total_amount", "total_due", "total", "amount_due", "balance_due", "grand_total", "amount"],
    "subtotal": ["subtotal", "sub_total", "net_amount", "sub_total_amount"],
    "tax": ["tax", "vat", "gst", "sales_tax", "tax_amount"],
    "currency": ["currency", "curr"],
    "supersedes": ["supersedes", "replaces", "revised_from"],
    "payment_terms": ["payment_terms", "terms"],
    "expiry_date": ["expiry_date", "expires", "expiration_date", "valid_until"],
}

_DATE_VALUE_PAT = (
    r"([0-9]{4}[-/][0-9]{1,2}[-/][0-9]{1,2}"
    r"|[0-9]{1,2}[-/][0-9]{1,2}[-/][0-9]{2,4}"
    r"|[A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4}"
    r"|\d{1,2}\s+[A-Za-z]{3,9}\s+\d{4})"
)

_CURR_PREFIX_PAT = r"(?:(?:INR|USD|EUR|GBP|CAD|AUD|JPY|CHF|SGD|NZD|CNY|Rs\.?)\s*)?"
_AMOUNT_TOKEN_PAT = r"([₹$€£¥]?\s*[\d,OoxX]+\.[\dOoxX]{2}\b|[₹$€£¥]?\s*[\d,]+\.\d{2})"

REGEX_PATTERNS: Dict[str, List[str]] = {
    "invoice_number": [
        r"(?i)\b(?:invoice\s*number|invoice\s*num|invoice\s*no\b|invoice\s*id|invoice\s*#|inv\s*#|inv#|inv\s*no\b|inv\s*num)\s*[:#\-]?\s*([A-Za-z0-9\-/_]+)",
        r"(?i)\b(?:invoice\s*code|bill\s*no\b|bill\s*number|reference\s*(?:no\b|number|#))\s*[:#\-]?\s*([A-Za-z0-9\-/_]+)",
        r"(?i)\b(?:invoice)\s*[:#]\s*([A-Za-z0-9\-/_]+)",
    ],
    "invoice_date": [
        rf"(?i)\b(?:invoice\s*date|bill\s*date|issue\s*date|date\s*of\s*issue|(?<!due\s)(?<!effective\s)date)\s*[:\-]?\s*{_DATE_VALUE_PAT}",
    ],
    "due_date": [
        rf"(?i)\b(?:due\s*date|payment\s*due|pay\s*by)\s*[:\-]?\s*{_DATE_VALUE_PAT}",
    ],
    "expiry_date": [
        rf"(?i)\b(?:expires?\s*on|expiry\s*date|expiration\s*date|continue\s*until|expiring\s*on)\s*[:\-]?\s*{_DATE_VALUE_PAT}",
    ],
    "total_amount": [
        rf"(?i)\b(?:total\s*due|total\s*amount|amount\s*due|balance\s*due|grand\s*total|(?<!sub)total)\s*[:\-]?\s*{_CURR_PREFIX_PAT}{_AMOUNT_TOKEN_PAT}",
    ],
    "subtotal": [
        rf"(?i)\b(?:subtotal|sub-total|sub\s*total|net\s*amount)\s*[:\-]?\s*{_CURR_PREFIX_PAT}{_AMOUNT_TOKEN_PAT}",
    ],
    "tax": [
        rf"(?i)\b(?:tax|vat|gst(?:\s*\(\d+%?\))?|sales\s*tax)\s*[:\-]?\s*{_CURR_PREFIX_PAT}{_AMOUNT_TOKEN_PAT}",
    ],
    "vendor": [
        r"(?i)\b(?:vendor|company|seller|supplier|billed\s+from|issued\s+by)\s*[:\-]\s*([A-Za-z0-9\s.,&'-]{2,60}?)(?=\s+\b(?:invoice|bill|date|due|subtotal|tax|total|currency)\b|$)",
        r"(?i)^([A-Za-z0-9\s.,&'-]{2,60}?)\s+(?:invoice|bill|statement)\b",
    ],
    "currency": [
        r"(?i)\b(INR|USD|EUR|GBP|CAD|AUD|JPY|CHF|SGD|NZD|CNY)\b",
    ],
    "supersedes": [
        r"(?i)\bsupersedes\s+(?:invoice\s+)?([A-Za-z0-9\-/_]+)",
    ],
    "payment_terms": [
        r"(?i)\bpayment\s*terms\s*[:\-]\s*([^\n\r.]+)",
    ],
}

CURRENCY_SYMBOL_MAP = {
    "₹": "INR",
    "$": "USD",
    "€": "EUR",
    "£": "GBP",
    "¥": "JPY",
}


def _clean_amount_string(raw: str) -> str:
    """Strip currency symbols and standard thousand commas while preserving corrupted OCR characters for verifier rejection."""
    s = raw.strip()
    s = re.sub(r"[₹$€£¥\s]", "", s)
    # If letters exist (e.g. '2O,0O0.5x'), return raw so verification flags corruption
    if re.search(r"[A-Za-z]", s):
        return raw.strip()
    if "." in s and "," in s:
        if s.rfind(",") > s.rfind("."):
            return s.replace(".", "").replace(",", ".")
        return s.replace(",", "")
    return s.replace(",", "")


def _detect_currency_from_text(text: str) -> Optional[Tuple[str, str]]:
    """Detect currency ISO code or symbol and return (currency_code, evidence_line)."""
    for line in text.splitlines():
        iso_match = re.search(r"\b(INR|USD|EUR|GBP|CAD|AUD|JPY|CHF|SGD|NZD|CNY)\b", line, re.IGNORECASE)
        if iso_match:
            return iso_match.group(1).upper(), line.strip()

    for line in text.splitlines():
        for sym, code in CURRENCY_SYMBOL_MAP.items():
            if sym in line:
                return code, line.strip()
    return None


def extract_from_json_dict(
    data: Dict[str, Any],
    requested_field: str,
    doc_path: str,
) -> Optional[ExtractedField]:
    """Extract a field from structured JSON data."""
    synonyms = FIELD_KEY_SYNONYMS.get(requested_field, [requested_field])

    for key, value in data.items():
        key_norm = key.lower().replace(" ", "_").replace("-", "_")
        if key_norm in synonyms:
            evidence = f'"{key}": {json.dumps(value)}'
            return ExtractedField(
                field_name=requested_field,
                extracted_value=value,
                source_document=doc_path,
                source_location=SourceLocation(description=f"JSON key '{key}'"),
                source_evidence=evidence,
                extraction_method="json_key",
            )

    for key, value in data.items():
        if isinstance(value, dict):
            nested_result = extract_from_json_dict(value, requested_field, doc_path)
            if nested_result:
                return nested_result
    return None


def extract_from_tables(
    tables: List[Dict[str, Any]],
    requested_field: str,
    doc_path: str,
) -> Optional[ExtractedField]:
    """Extract a field from tabular structures (CSV / XLSX) with accurate cell coordinates."""
    synonyms = FIELD_KEY_SYNONYMS.get(requested_field, [requested_field])

    for tbl in tables:
        sheet_name = tbl.get("sheet")
        headers = tbl.get("headers", [])
        rows = tbl.get("rows", [])
        all_rows = tbl.get("all_rows") or (([headers] if headers else []) + rows)
        all_coords = tbl.get("all_coords")

        # 1. Key-value style: check all rows where col_idx has label and next non-empty cell has value
        for row_idx, row in enumerate(all_rows, start=1):
            if len(row) >= 2 and row[0] is not None:
                col0 = str(row[0]).strip().lower().replace(":", "").replace(" ", "_").replace("-", "_")
                if any(syn == col0 or col0.startswith(syn) for syn in synonyms):
                    # Find first non-empty value column after col 0
                    for val_col_idx in range(1, len(row)):
                        if row[val_col_idx] is not None and str(row[val_col_idx]).strip():
                            val = str(row[val_col_idx]).strip()
                            if all_coords and row_idx - 1 < len(all_coords) and val_col_idx < len(all_coords[row_idx - 1]):
                                cell_coord = all_coords[row_idx - 1][val_col_idx]
                            else:
                                col_letter = chr(ord("A") + min(val_col_idx, 25))
                                cell_coord = f"{col_letter}{row_idx}" if sheet_name else None

                            evidence = f"{row[0]} = {row[val_col_idx]}"
                            loc = SourceLocation(
                                sheet=sheet_name,
                                cell=cell_coord,
                                line_number=row_idx,
                                description=f"Row {row_idx}" if not sheet_name else None,
                            )
                            return ExtractedField(
                                field_name=requested_field,
                                extracted_value=_clean_amount_string(val) if requested_field in ("total_amount", "subtotal", "tax") else val,
                                source_document=doc_path,
                                source_location=loc,
                                source_evidence=evidence,
                                extraction_method="tabular_cell",
                            )

        # 2. Column-header style
        matched_col = None
        for col_idx, h in enumerate(headers):
            h_norm = str(h).strip().lower().replace(" ", "_").replace(":", "")
            if any(syn == h_norm for syn in synonyms):
                matched_col = col_idx
                break

        if matched_col is not None and rows:
            candidate_rows = reversed(list(enumerate(rows, start=2))) if requested_field in ("total_amount", "subtotal", "tax") else enumerate(rows, start=2)
            for row_idx, row in candidate_rows:
                if matched_col < len(row) and row[matched_col] is not None:
                    val = str(row[matched_col]).strip()
                    if val:
                        col_letter = chr(ord("A") + min(matched_col, 25))
                        evidence = f"{headers[matched_col]}: {val}"
                        loc = SourceLocation(
                            sheet=sheet_name,
                            cell=f"{col_letter}{row_idx}" if sheet_name else None,
                            line_number=row_idx,
                            description=f"Column '{headers[matched_col]}'",
                        )
                        return ExtractedField(
                            field_name=requested_field,
                            extracted_value=_clean_amount_string(val) if requested_field in ("total_amount", "subtotal", "tax") else val,
                            source_document=doc_path,
                            source_location=loc,
                            source_evidence=evidence,
                            extraction_method="tabular_cell",
                        )
    return None


def _is_buyer_or_address_line(line: str) -> bool:
    """Return True if a line represents a Bill-To/Buyer line, metadata header, or street address rather than a vendor name."""
    low = line.lower().strip()
    if re.search(
        r"^(?:bill\s*to|billed\s*to|ship\s*to|buyer|customer|invoice\s*number|invoice\s*date|due\s*date|payment\s*terms|revised\s*invoice|supersedes|tax\s*invoice|revised\b)\b",
        low,
    ):
        return True
    # Street address lines starting with numbers or Plot
    if re.match(r"^(?:\d+\s+|plot\s+\d+)", low):
        return True
    return False


def extract_from_pages(
    pages: List[DocumentPage],
    requested_field: str,
    doc_path: str,
) -> Optional[ExtractedField]:
    """Extract a field from text pages using regex and line grounding."""
    if requested_field == "currency":
        for page in pages:
            detected = _detect_currency_from_text(page.text)
            if detected:
                code, ev_line = detected
                loc = SourceLocation(page=page.page_num, description=f"Page {page.page_num}")
                method = "ocr_regex" if page.is_ocr else "text_regex"
                return ExtractedField(
                    field_name="currency",
                    extracted_value=code,
                    source_document=doc_path,
                    source_location=loc,
                    source_evidence=ev_line,
                    extraction_method=method,
                )

    patterns = list(REGEX_PATTERNS.get(requested_field, []))
    if not patterns:
        # Dynamic fallback pattern for custom caller-requested fields (e.g. "po_number")
        label_pat = requested_field.replace("_", r"[\s_\-]+")
        patterns = [rf"(?i)\b{label_pat}\s*[:\-]\s*([^\n\r]+)"]

    for page in pages:
        lines = page.text.splitlines()
        method = "ocr_regex" if page.is_ocr else "text_regex"

        if requested_field == "vendor":
            # Pass 1: Explicit labeled vendor lines (e.g. "Vendor: Acme Corp")
            for line_idx, line in enumerate(lines, start=1):
                clean_line = line.strip()
                if not clean_line or clean_line.startswith("---") or _is_buyer_or_address_line(clean_line):
                    continue
                for pat in patterns:
                    m = re.search(pat, clean_line)
                    if m:
                        val = m.group(1).strip()
                        if not _is_buyer_or_address_line(val):
                            return ExtractedField(
                                field_name=requested_field,
                                extracted_value=val,
                                source_document=doc_path,
                                source_location=SourceLocation(page=page.page_num, line_number=line_idx),
                                source_evidence=clean_line,
                                extraction_method=method,
                            )

            # Pass 2: Line immediately after standalone "INVOICE" title header
            for line_idx, line in enumerate(lines, start=1):
                clean_line = line.strip()
                if clean_line.upper() == "INVOICE" and line_idx < len(lines):
                    next_line = lines[line_idx].strip()
                    if next_line and not _is_buyer_or_address_line(next_line):
                        return ExtractedField(
                            field_name=requested_field,
                            extracted_value=next_line,
                            source_document=doc_path,
                            source_location=SourceLocation(page=page.page_num, line_number=line_idx + 1),
                            source_evidence=next_line,
                            extraction_method=method,
                        )

            # Pass 3: Corporate entity line (excluding Bill-To / Buyer / Address lines)
            corp_keywords = (
                "inc", "corp", "corporation", "llc", "ltd", "supplies", "logistics",
                "components", "chemicals", "services", "systems", "technologies", "enterprises",
            )
            for line_idx, line in enumerate(lines, start=1):
                clean_line = line.strip()
                if not clean_line or clean_line.startswith("---") or _is_buyer_or_address_line(clean_line):
                    continue
                if any(re.search(rf"\b{kw}\b", clean_line.lower()) for kw in corp_keywords):
                    return ExtractedField(
                        field_name=requested_field,
                        extracted_value=clean_line,
                        source_document=doc_path,
                        source_location=SourceLocation(page=page.page_num, line_number=line_idx),
                        source_evidence=clean_line,
                        extraction_method=method,
                    )

        # Standard line-by-line regex search (including 2-line label/value pairs from columnar PDF text)
        for line_idx, line in enumerate(lines, start=1):
            candidates_to_check = [line]
            if line.strip().endswith(":") and line_idx < len(lines):
                candidates_to_check.append(f"{line.strip()} {lines[line_idx].strip()}")
            for candidate_line in candidates_to_check:
                for pat in patterns:
                    m = re.search(pat, candidate_line)
                    if m:
                        raw_val = m.group(1).strip()
                        val = (
                            _clean_amount_string(raw_val)
                            if requested_field in ("total_amount", "subtotal", "tax")
                            else raw_val
                        )
                        return ExtractedField(
                            field_name=requested_field,
                            extracted_value=val,
                            source_document=doc_path,
                            source_location=SourceLocation(
                                page=page.page_num,
                                line_number=line_idx,
                                span=m.span(1),
                            ),
                            source_evidence=candidate_line.strip(),
                            extraction_method=method,
                        )
    return None


def extract_fields(
    doc: DocumentContent,
    requested_fields: Optional[List[str]] = None,
) -> ExtractionResult:
    """Extract caller-specified fields from document content preserving source evidence."""
    target_fields = requested_fields or DEFAULT_REQUESTED_FIELDS
    doc_path_str = str(doc.file_path)

    result = ExtractionResult(
        document_path=doc_path_str,
        candidate_documents=[doc_path_str],
        recovery_log=list(doc.metadata.get("recovery_log", [])),
    )

    for field_name in target_fields:
        field_extracted: Optional[ExtractedField] = None

        if doc.structured_data and isinstance(doc.structured_data, dict):
            field_extracted = extract_from_json_dict(doc.structured_data, field_name, doc_path_str)

        if not field_extracted and doc.tables:
            field_extracted = extract_from_tables(doc.tables, field_name, doc_path_str)

        if not field_extracted and doc.pages:
            field_extracted = extract_from_pages(doc.pages, field_name, doc_path_str)

        if field_extracted and field_extracted.extracted_value is not None:
            result.fields[field_name] = field_extracted
        else:
            missing_field = ExtractedField(
                field_name=field_name,
                extracted_value=None,
                source_document=doc_path_str,
                unresolved_issue=f"Field '{field_name}' not found in document",
            )
            result.fields[field_name] = missing_field
            result.add_failure(
                category=FailureCategory.MISSING_INFORMATION,
                message=f"Requested field '{field_name}' could not be located",
                field_name=field_name,
                file_path=doc_path_str,
            )

    return result


def process_invoice_request(
    data_dir: str | Path,
    query: str,
    requested_fields: Optional[List[str]] = None,
    ocr_engine: Optional[OCREngine] = None,
) -> ExtractionResult:
    """Find, read, extract, and verify information from a relevant invoice in data_dir."""
    dir_path = Path(data_dir)
    candidates, is_ambiguous, reason = find_candidate_documents(dir_path, query, ocr_engine=ocr_engine)

    if not candidates:
        res = ExtractionResult(
            status="failed",
            requires_human_review=True,
            summary_notes=[reason or f"No invoice or document found for query: {query}"],
        )
        res.add_failure(
            category=FailureCategory.MISSING_INFORMATION,
            message=reason or f"No candidate document found for query: {query}",
        )
        return res

    if is_ambiguous:
        res = ExtractionResult(
            candidate_documents=[str(p) for p in candidates],
            status="ambiguous",
            requires_human_review=True,
            summary_notes=[reason or "Ambiguous candidate invoices"],
        )
        res.add_failure(
            category=FailureCategory.AMBIGUOUS_OR_CONFLICTING_INFORMATION,
            message=reason or f"Multiple candidate documents found for query: {query}",
            details={"candidates": [str(p) for p in candidates]},
        )
        return res

    candidate_file = candidates[0]
    try:
        doc = read_document(candidate_file, ocr_engine=ocr_engine)
    except Exception as exc:
        cat = classify_exception(exc)
        res = ExtractionResult(
            document_path=str(candidate_file),
            candidate_documents=[str(candidate_file)],
            status="failed",
            requires_human_review=True,
            summary_notes=[f"Failed reading document: {exc}"],
        )
        res.add_failure(
            category=cat,
            message=str(exc),
            file_path=str(candidate_file),
        )
        return res

    result = extract_fields(doc, requested_fields=requested_fields)
    from agents.verification import verify_extraction_result
    result = verify_extraction_result(result, doc_raw_text=doc.raw_text)
    return result
